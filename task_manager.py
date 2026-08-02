from flask import Flask, render_template, request, jsonify, Response
import sqlite3
import calendar
import csv
import io
import json
from datetime import datetime, timedelta
import os
from dotenv import load_dotenv
load_dotenv()
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_XLSX = True
except ImportError:
    HAS_XLSX = False

app = Flask(__name__)
DB_PATH = os.environ.get("DB_PATH", "tasks.db")


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db():
    with get_db() as conn:
        conn.executescript("""
            CREATE TABLE IF NOT EXISTS tasks (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT NOT NULL,
                description TEXT DEFAULT '',
                status TEXT NOT NULL DEFAULT 'not_started',
                priority TEXT NOT NULL DEFAULT 'medium',
                due_date TEXT DEFAULT NULL,
                due_time TEXT DEFAULT NULL,
                estimated_minutes INTEGER DEFAULT 0,
                project TEXT DEFAULT NULL,
                recurrence TEXT NOT NULL DEFAULT 'none',
                task_type TEXT NOT NULL DEFAULT 'general',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS subtasks (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                task_id INTEGER NOT NULL,
                title TEXT NOT NULL,
                is_done INTEGER NOT NULL DEFAULT 0,
                position INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL,
                FOREIGN KEY (task_id) REFERENCES tasks(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS tags (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL UNIQUE,
                color TEXT NOT NULL DEFAULT '#3b82f6'
            );
            CREATE TABLE IF NOT EXISTS task_tags (
                task_id INTEGER NOT NULL,
                tag_id INTEGER NOT NULL,
                PRIMARY KEY (task_id, tag_id),
                FOREIGN KEY (task_id) REFERENCES tasks(id) ON DELETE CASCADE,
                FOREIGN KEY (tag_id) REFERENCES tags(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS activity_log (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                task_id INTEGER NOT NULL,
                action TEXT NOT NULL,
                detail TEXT DEFAULT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY (task_id) REFERENCES tasks(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS time_entries (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                task_id INTEGER NOT NULL,
                started_at TEXT NOT NULL,
                ended_at TEXT DEFAULT NULL,
                duration_seconds INTEGER DEFAULT NULL,
                FOREIGN KEY (task_id) REFERENCES tasks(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS reminders (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                task_id INTEGER NOT NULL,
                remind_at TEXT NOT NULL,
                recurrence TEXT NOT NULL DEFAULT 'none',
                sound INTEGER NOT NULL DEFAULT 1,
                is_active INTEGER NOT NULL DEFAULT 1,
                last_triggered TEXT DEFAULT NULL,
                snoozed_until TEXT DEFAULT NULL,
                FOREIGN KEY (task_id) REFERENCES tasks(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS meetings (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT NOT NULL,
                description TEXT DEFAULT '',
                date TEXT NOT NULL,
                time TEXT DEFAULT NULL,
                duration_minutes INTEGER DEFAULT 60,
                location TEXT DEFAULT '',
                recurrence TEXT NOT NULL DEFAULT 'none',
                status TEXT NOT NULL DEFAULT 'upcoming',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS meeting_points (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                meeting_id INTEGER NOT NULL,
                content TEXT NOT NULL,
                is_done INTEGER NOT NULL DEFAULT 0,
                position INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL,
                FOREIGN KEY (meeting_id) REFERENCES meetings(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS standalone_reminders (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT NOT NULL,
                note TEXT DEFAULT '',
                reminder_type TEXT NOT NULL DEFAULT 'general',
                remind_at TEXT NOT NULL,
                recurrence TEXT NOT NULL DEFAULT 'none',
                sound INTEGER NOT NULL DEFAULT 1,
                is_active INTEGER NOT NULL DEFAULT 1,
                last_triggered TEXT DEFAULT NULL,
                snoozed_until TEXT DEFAULT NULL,
                created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS task_steps (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                task_id INTEGER NOT NULL,
                title TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'not_started',
                position INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                FOREIGN KEY (task_id) REFERENCES tasks(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS task_dependencies (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                task_id INTEGER NOT NULL,
                depends_on_id INTEGER NOT NULL,
                created_at TEXT NOT NULL,
                UNIQUE(task_id, depends_on_id),
                FOREIGN KEY (task_id) REFERENCES tasks(id) ON DELETE CASCADE,
                FOREIGN KEY (depends_on_id) REFERENCES tasks(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS note_books (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT NOT NULL,
                color TEXT NOT NULL DEFAULT 'default',
                position INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS note_chapters (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                book_id INTEGER NOT NULL,
                title TEXT NOT NULL,
                position INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                FOREIGN KEY (book_id) REFERENCES note_books(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS note_pages (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                chapter_id INTEGER NOT NULL,
                title TEXT NOT NULL DEFAULT '',
                body TEXT NOT NULL DEFAULT '',
                color TEXT NOT NULL DEFAULT 'default',
                pinned INTEGER NOT NULL DEFAULT 0,
                position INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                FOREIGN KEY (chapter_id) REFERENCES note_chapters(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS ai_settings (
                id INTEGER PRIMARY KEY CHECK (id = 1),
                base_url TEXT NOT NULL DEFAULT 'http://localhost:11434',
                model TEXT NOT NULL DEFAULT 'qwen3:8b',
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS ai_conversations (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                context_type TEXT NOT NULL,
                context_id INTEGER NOT NULL,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                UNIQUE(context_type, context_id)
            );
            CREATE TABLE IF NOT EXISTS ai_messages (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                conversation_id INTEGER NOT NULL,
                role TEXT NOT NULL,
                mode TEXT DEFAULT NULL,
                content TEXT NOT NULL,
                questions TEXT DEFAULT NULL,
                sections TEXT DEFAULT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY (conversation_id) REFERENCES ai_conversations(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS ai_suggestions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                conversation_id INTEGER NOT NULL,
                message_id INTEGER DEFAULT NULL,
                title TEXT NOT NULL,
                reasoning TEXT DEFAULT '',
                kind TEXT NOT NULL DEFAULT 'step',
                status TEXT NOT NULL DEFAULT 'pending',
                created_task_id INTEGER DEFAULT NULL,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                FOREIGN KEY (conversation_id) REFERENCES ai_conversations(id) ON DELETE CASCADE,
                FOREIGN KEY (message_id) REFERENCES ai_messages(id) ON DELETE SET NULL
            );
        """)
        if not conn.execute("SELECT id FROM ai_settings WHERE id=1").fetchone():
            conn.execute(
                "INSERT INTO ai_settings (id,base_url,model,updated_at) VALUES (1,?,?,?)",
                ("http://localhost:11434", "qwen3:8b", _now()),
            )
        for col, defn in [
            ("priority",           "TEXT NOT NULL DEFAULT 'medium'"),
            ("due_date",           "TEXT DEFAULT NULL"),
            ("due_time",           "TEXT DEFAULT NULL"),
            ("estimated_minutes",  "INTEGER DEFAULT 0"),
            ("project",            "TEXT DEFAULT NULL"),
            ("recurrence",         "TEXT NOT NULL DEFAULT 'none'"),
            ("task_type",          "TEXT NOT NULL DEFAULT 'general'"),
            ("step_id",            "INTEGER DEFAULT NULL"),
        ]:
            try:
                conn.execute(f"ALTER TABLE tasks ADD COLUMN {col} {defn}")
            except Exception:
                pass
        for col, defn in [("snoozed_until", "TEXT DEFAULT NULL")]:
            try:
                conn.execute(f"ALTER TABLE reminders ADD COLUMN {col} {defn}")
            except Exception:
                pass
        for col, defn in [("parallel_group", "TEXT DEFAULT NULL")]:
            try:
                conn.execute(f"ALTER TABLE task_steps ADD COLUMN {col} {defn}")
            except Exception:
                pass
        for col, defn in [
            ("provider",        "TEXT NOT NULL DEFAULT 'ollama'"),
            ("anthropic_model", "TEXT NOT NULL DEFAULT 'claude-sonnet-5'"),
        ]:
            try:
                conn.execute(f"ALTER TABLE ai_settings ADD COLUMN {col} {defn}")
            except Exception:
                pass
        for col, defn in [
            ("patch_action", "TEXT DEFAULT NULL"),
            ("patch_html",   "TEXT DEFAULT NULL"),
            ("patch_find",   "TEXT DEFAULT NULL"),
        ]:
            try:
                conn.execute(f"ALTER TABLE ai_suggestions ADD COLUMN {col} {defn}")
            except Exception:
                pass
        conn.commit()
        _migrate_step_points_and_milestones(conn)


def _migrate_step_points_and_milestones(conn):
    """One-time migration: fold step_points into real tasks (step_id-linked),
    then drop the now-unused step_points and milestones tables/columns."""
    tables = {r["name"] for r in conn.execute(
        "SELECT name FROM sqlite_master WHERE type='table'"
    ).fetchall()}

    if "step_points" in tables:
        now = _now()
        status_map = {"not_started": "not_started", "in_process": "in_process",
                      "completed": "completed", "cancelled": "cancelled"}
        for p in conn.execute("SELECT * FROM step_points").fetchall():
            desc_parts = []
            if p["notes"]:
                desc_parts.append(f"Notes: {p['notes']}")
            if p["blocker"]:
                desc_parts.append(f"Blocker: {p['blocker']}")
            if p["solution"]:
                desc_parts.append(f"Solution: {p['solution']}")
            conn.execute(
                "INSERT INTO tasks (title,description,status,priority,task_type,step_id,recurrence,created_at,updated_at)"
                " VALUES (?,?,?,?,?,?,?,?,?)",
                (p["content"], "\n".join(desc_parts), status_map.get(p["status"], "not_started"),
                 "medium", "general", p["step_id"], "none", p["created_at"], now),
            )
        conn.execute("DROP TABLE step_points")
        conn.commit()

    if "milestones" in tables:
        conn.execute("DROP TABLE milestones")
        conn.commit()

    cols = {r["name"] for r in conn.execute("PRAGMA table_info(tasks)").fetchall()}
    if "milestone_id" in cols:
        conn.execute("ALTER TABLE tasks DROP COLUMN milestone_id")
        conn.commit()

    # Flashcards module was removed entirely (superseded by the AI Assistant) --
    # drop its tables and any data with them.
    if "flashcards" in tables:
        conn.execute("DROP TABLE flashcards")
        conn.commit()
    if "flashcard_categories" in tables:
        conn.execute("DROP TABLE flashcard_categories")
        conn.commit()


# ── Helpers ───────────────────────────────────────────────────────────────────

def _now():
    return datetime.now().isoformat(timespec="seconds")


def _log(conn, task_id, action, detail=None):
    conn.execute(
        "INSERT INTO activity_log (task_id, action, detail, created_at) VALUES (?,?,?,?)",
        (task_id, action, detail, _now()),
    )


def _add_months(dt, n):
    month = dt.month - 1 + n
    year  = dt.year + month // 12
    month = month % 12 + 1
    day   = min(dt.day, calendar.monthrange(year, month)[1])
    return dt.replace(year=year, month=month, day=day)


def _next_occurrence(remind_at_iso, recurrence):
    dt  = datetime.fromisoformat(remind_at_iso)
    now = datetime.now()
    if recurrence == "daily":
        while dt <= now:
            dt += timedelta(days=1)
    elif recurrence == "weekdays":
        dt += timedelta(days=1)
        while dt <= now or dt.weekday() >= 5:
            dt += timedelta(days=1)
    elif recurrence == "weekly":
        while dt <= now:
            dt += timedelta(weeks=1)
    elif recurrence == "monthly":
        while dt <= now:
            dt = _add_months(dt, 1)
    return dt.isoformat(timespec="seconds")


def _next_task_due(due_date_str, recurrence):
    """Return the next due date string for a recurring task after completion."""
    try:
        base = datetime.strptime(due_date_str, "%Y-%m-%d")
    except (ValueError, TypeError):
        base = datetime.now()
    if recurrence == "daily":
        nxt = base + timedelta(days=1)
    elif recurrence == "weekdays":
        nxt = base + timedelta(days=1)
        while nxt.weekday() >= 5:   # skip Sat(5) and Sun(6)
            nxt += timedelta(days=1)
    elif recurrence == "weekly":
        nxt = base + timedelta(weeks=1)
    elif recurrence == "biweekly":
        nxt = base + timedelta(weeks=2)
    elif recurrence == "monthly":
        nxt = _add_months(base, 1)
    else:
        return None
    return nxt.strftime("%Y-%m-%d")


def _spawn_next(conn, task, now):
    """Create the next occurrence of a recurring task and copy its tags."""
    next_due = _next_task_due(task["due_date"], task["recurrence"])
    if not next_due:
        return
    cur = conn.execute(
        "INSERT INTO tasks (title,description,status,priority,due_date,due_time,"
        "estimated_minutes,project,recurrence,task_type,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
        (task["title"], task["description"], "not_started", task["priority"],
         next_due, task["due_time"], task["estimated_minutes"] or 0,
         task["project"], task["recurrence"], dict(task).get("task_type","general"), now, now),
    )
    new_id = cur.lastrowid
    # copy tags
    tag_ids = conn.execute(
        "SELECT tag_id FROM task_tags WHERE task_id=?", (task["id"],)
    ).fetchall()
    for row in tag_ids:
        conn.execute("INSERT OR IGNORE INTO task_tags (task_id,tag_id) VALUES (?,?)",
                     (new_id, row["tag_id"]))
    _log(conn, new_id, "created", f"Auto-spawned from #{task['id']} ({task['recurrence']})")


def _enrich_meetings(conn, meetings):
    result = []
    for m in meetings:
        mt = dict(m)
        mid = mt["id"]
        pts = conn.execute(
            "SELECT COUNT(*) total, SUM(is_done) done FROM meeting_points WHERE meeting_id=?", (mid,)
        ).fetchone()
        mt["point_total"] = pts["total"] or 0
        mt["point_done"]  = int(pts["done"] or 0)
        preview = conn.execute(
            "SELECT content FROM meeting_points WHERE meeting_id=? AND is_done=0 ORDER BY position, id LIMIT 3",
            (mid,)
        ).fetchall()
        mt["point_preview"] = [p["content"] for p in preview]
        result.append(mt)
    return result


def _spawn_next_meeting(conn, m, now):
    next_date = _next_task_due(m["date"], m["recurrence"])
    if not next_date:
        return
    cur = conn.execute(
        "INSERT INTO meetings (title,description,date,time,duration_minutes,location,recurrence,status,created_at,updated_at)"
        " VALUES (?,?,?,?,?,?,?,?,?,?)",
        (m["title"], m["description"], next_date, m["time"],
         m["duration_minutes"], m["location"], m["recurrence"], "upcoming", now, now),
    )
    new_id = cur.lastrowid
    # Copy discussion points (all unchecked for next occurrence)
    for p in conn.execute(
        "SELECT * FROM meeting_points WHERE meeting_id=? ORDER BY position", (m["id"],)
    ).fetchall():
        conn.execute(
            "INSERT INTO meeting_points (meeting_id,content,is_done,position,created_at) VALUES (?,?,0,?,?)",
            (new_id, p["content"], p["position"], now),
        )


def _enrich(conn, tasks):
    """Attach subtask counts, tags, time totals, and active timer to each task dict."""
    result = []
    for task in tasks:
        t   = dict(task)
        tid = t["id"]

        sub = conn.execute(
            "SELECT COUNT(*) total, SUM(is_done) done FROM subtasks WHERE task_id=?", (tid,)
        ).fetchone()
        t["subtask_total"] = sub["total"] or 0
        t["subtask_done"]  = int(sub["done"] or 0)

        sub_preview = conn.execute(
            "SELECT title FROM subtasks WHERE task_id=? AND is_done=0 ORDER BY position, id LIMIT 3",
            (tid,)
        ).fetchall()
        t["subtask_preview"] = [s["title"] for s in sub_preview]

        tags = conn.execute(
            "SELECT tg.id, tg.name, tg.color FROM tags tg "
            "JOIN task_tags tt ON tg.id=tt.tag_id WHERE tt.task_id=?", (tid,)
        ).fetchall()
        t["tags"] = [dict(tg) for tg in tags]

        te = conn.execute(
            "SELECT SUM(duration_seconds) total, "
            "MAX(CASE WHEN ended_at IS NULL THEN started_at END) active "
            "FROM time_entries WHERE task_id=?", (tid,)
        ).fetchone()
        t["time_total"]         = int(te["total"] or 0)
        t["active_timer_since"] = te["active"]

        rem = conn.execute(
            "SELECT id, remind_at, recurrence FROM reminders "
            "WHERE task_id=? AND is_active=1 LIMIT 1", (tid,)
        ).fetchone()
        t["reminder"] = dict(rem) if rem else None

        step_info = conn.execute(
            "SELECT COUNT(*) total, "
            "SUM(CASE WHEN status='completed' THEN 1 ELSE 0 END) done "
            "FROM task_steps WHERE task_id=?", (tid,)
        ).fetchone()
        t["step_total"] = step_info["total"] or 0
        t["step_done"]  = int(step_info["done"] or 0)

        step = None
        if t.get("step_id"):
            srow = conn.execute(
                "SELECT ts.id, ts.title, ts.status, ts.task_id AS workflow_id, wf.title AS workflow_title "
                "FROM task_steps ts JOIN tasks wf ON wf.id=ts.task_id WHERE ts.id=?", (t["step_id"],)
            ).fetchone()
            step = dict(srow) if srow else None
        t["step"] = step

        depends_on = [dict(r) for r in conn.execute(
            "SELECT tk.id, tk.title, tk.status FROM task_dependencies d "
            "JOIN tasks tk ON tk.id=d.depends_on_id WHERE d.task_id=?", (tid,)
        ).fetchall()]
        blocks = [dict(r) for r in conn.execute(
            "SELECT tk.id, tk.title, tk.status FROM task_dependencies d "
            "JOIN tasks tk ON tk.id=d.task_id WHERE d.depends_on_id=?", (tid,)
        ).fetchall()]
        t["depends_on"] = depends_on
        t["blocks"]     = blocks
        t["is_blocked"] = any(d["status"] not in ("completed", "cancelled") for d in depends_on)

        result.append(t)
    return result


def _has_dependency_path(conn, start_id, target_id, visited=None):
    """True if start_id depends on target_id, directly or transitively."""
    if visited is None:
        visited = set()
    if start_id in visited:
        return False
    visited.add(start_id)
    if start_id == target_id:
        return True
    rows = conn.execute(
        "SELECT depends_on_id FROM task_dependencies WHERE task_id=?", (start_id,)
    ).fetchall()
    return any(_has_dependency_path(conn, r["depends_on_id"], target_id, visited) for r in rows)


# ── Index ─────────────────────────────────────────────────────────────────────

@app.route("/")
def index():
    resp = Response(render_template("task_manager.html"))
    resp.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    return resp


# ── Tasks ─────────────────────────────────────────────────────────────────────

@app.route("/tasks")
def get_tasks():
    with get_db() as conn:
        rows = conn.execute("SELECT * FROM tasks ORDER BY updated_at DESC").fetchall()
        return jsonify(_enrich(conn, rows))


@app.route("/tasks/export.csv")
def export_csv():
    type_filter = request.args.get("type", "all")
    with get_db() as conn:
        if type_filter and type_filter != "all":
            rows = conn.execute(
                "SELECT * FROM tasks WHERE task_type=? ORDER BY created_at", (type_filter,)
            ).fetchall()
        else:
            rows = conn.execute("SELECT * FROM tasks ORDER BY created_at").fetchall()
        tasks = _enrich(conn, rows)
    buf = io.StringIO()
    w   = csv.writer(buf)
    w.writerow(["ID","Title","Type","Description","Status","Priority","Due Date","Due Time",
                "Est. (min)","Project","Workflow","Step","Recurrence","Tags","Subtasks","Time (hrs)","Created","Updated"])
    for t in tasks:
        step = t.get("step")
        w.writerow([
            t["id"], t["title"], dict(t).get("task_type","general"),
            t["description"], t["status"], t["priority"],
            t["due_date"] or "", t["due_time"] or "", t["estimated_minutes"] or 0,
            t["project"] or "", step["workflow_title"] if step else "", step["title"] if step else "",
            t["recurrence"] or "none",
            ",".join(tg["name"] for tg in t["tags"]),
            f"{t['subtask_done']}/{t['subtask_total']}",
            round(t["time_total"] / 3600, 2),
            t["created_at"], t["updated_at"],
        ])
    fname = f"tasks_{type_filter}.csv" if type_filter != "all" else "tasks.csv"
    return Response(buf.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": f"attachment; filename={fname}"})


@app.route("/tasks", methods=["POST"])
def create_task():
    data  = request.get_json() or {}
    title = data.get("title", "").strip()
    if not title:
        return jsonify({"error": "Title is required"}), 400
    now = _now()
    with get_db() as conn:
        cur = conn.execute(
            "INSERT INTO tasks (title,description,status,priority,due_date,due_time,estimated_minutes,project,recurrence,task_type,step_id,created_at,updated_at)"
            " VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (title, data.get("description","").strip(), data.get("status","not_started"),
             data.get("priority","medium"), data.get("due_date") or None,
             data.get("due_time") or None, int(data.get("estimated_minutes") or 0),
             data.get("project") or None, data.get("recurrence","none") or "none",
             data.get("task_type","general") or "general",
             int(data["step_id"]) if data.get("step_id") else None, now, now),
        )
        tid = cur.lastrowid
        _log(conn, tid, "created", title)
        conn.commit()
        task = conn.execute("SELECT * FROM tasks WHERE id=?", (tid,)).fetchone()
        return jsonify(_enrich(conn, [task])[0]), 201


@app.route("/tasks/<int:tid>", methods=["PUT"])
def update_task(tid):
    data = request.get_json() or {}
    with get_db() as conn:
        task = conn.execute("SELECT * FROM tasks WHERE id=?", (tid,)).fetchone()
        if not task:
            return jsonify({"error": "Not found"}), 404
        now          = _now()
        task_d = dict(task)
        new_title    = data.get("title",    task_d["title"]).strip() or task_d["title"]
        new_status   = data.get("status",   task_d["status"])
        new_priority = data.get("priority", task_d["priority"])
        new_recurrence = data.get("recurrence", task_d["recurrence"]) or "none"
        new_task_type  = data.get("task_type",  task_d.get("task_type","general")) or "general"
        if new_status   != task_d["status"]:   _log(conn, tid, "status_changed",   f"{task_d['status']} → {new_status}")
        if new_priority != task_d["priority"]: _log(conn, tid, "priority_changed", f"{task_d['priority']} → {new_priority}")
        if new_title    != task_d["title"]:    _log(conn, tid, "title_changed",    f"→ {new_title}")
        new_step_id = data["step_id"] if "step_id" in data else task_d.get("step_id")
        new_step_id = int(new_step_id) if new_step_id else None
        conn.execute(
            "UPDATE tasks SET title=?,description=?,status=?,priority=?,due_date=?,due_time=?,estimated_minutes=?,project=?,recurrence=?,task_type=?,step_id=?,updated_at=? WHERE id=?",
            (new_title, data.get("description", task_d["description"]),
             new_status, new_priority,
             data.get("due_date", task_d["due_date"]) or None,
             data.get("due_time", task_d["due_time"]) or None,
             int(data.get("estimated_minutes", task_d["estimated_minutes"]) or 0),
             data.get("project",  task_d["project"])  or None,
             new_recurrence, new_task_type, new_step_id, now, tid),
        )
        # Auto-spawn next occurrence when a recurring task is completed
        if new_status == "completed" and task_d["status"] != "completed" and new_recurrence != "none":
            updated_task = conn.execute("SELECT * FROM tasks WHERE id=?", (tid,)).fetchone()
            _spawn_next(conn, updated_task, now)
        conn.commit()
        updated = conn.execute("SELECT * FROM tasks WHERE id=?", (tid,)).fetchone()
        return jsonify(_enrich(conn, [updated])[0])


@app.route("/tasks/<int:tid>", methods=["DELETE"])
def delete_task(tid):
    with get_db() as conn:
        if not conn.execute("SELECT id FROM tasks WHERE id=?", (tid,)).fetchone():
            return jsonify({"error": "Not found"}), 404
        conn.execute("DELETE FROM tasks WHERE id=?", (tid,))
        conn.commit()
        return jsonify({"success": True})


# ── Subtasks ──────────────────────────────────────────────────────────────────

@app.route("/tasks/<int:tid>/subtasks")
def get_subtasks(tid):
    with get_db() as conn:
        rows = conn.execute(
            "SELECT * FROM subtasks WHERE task_id=? ORDER BY position, id", (tid,)
        ).fetchall()
        return jsonify([dict(r) for r in rows])


@app.route("/tasks/<int:tid>/subtasks", methods=["POST"])
def add_subtask(tid):
    data  = request.get_json() or {}
    title = data.get("title", "").strip()
    if not title:
        return jsonify({"error": "Title required"}), 400
    now = _now()
    with get_db() as conn:
        pos = conn.execute(
            "SELECT COALESCE(MAX(position),0)+1 FROM subtasks WHERE task_id=?", (tid,)
        ).fetchone()[0]
        cur = conn.execute(
            "INSERT INTO subtasks (task_id,title,is_done,position,created_at) VALUES (?,?,0,?,?)",
            (tid, title, pos, now),
        )
        _log(conn, tid, "subtask_added", title)
        conn.commit()
        row = conn.execute("SELECT * FROM subtasks WHERE id=?", (cur.lastrowid,)).fetchone()
        return jsonify(dict(row)), 201


@app.route("/subtasks/<int:sid>", methods=["PUT"])
def update_subtask(sid):
    data = request.get_json() or {}
    with get_db() as conn:
        st = conn.execute("SELECT * FROM subtasks WHERE id=?", (sid,)).fetchone()
        if not st:
            return jsonify({"error": "Not found"}), 404
        is_done = data.get("is_done", st["is_done"])
        title   = data.get("title",   st["title"])
        if is_done != st["is_done"]:
            _log(conn, st["task_id"],
                 "subtask_completed" if is_done else "subtask_reopened", title)
        conn.execute("UPDATE subtasks SET title=?,is_done=? WHERE id=?", (title, is_done, sid))
        conn.execute("UPDATE tasks SET updated_at=? WHERE id=?", (_now(), st["task_id"]))
        conn.commit()
        return jsonify(dict(conn.execute("SELECT * FROM subtasks WHERE id=?", (sid,)).fetchone()))


@app.route("/subtasks/<int:sid>", methods=["DELETE"])
def delete_subtask(sid):
    with get_db() as conn:
        if not conn.execute("SELECT id FROM subtasks WHERE id=?", (sid,)).fetchone():
            return jsonify({"error": "Not found"}), 404
        conn.execute("DELETE FROM subtasks WHERE id=?", (sid,))
        conn.commit()
        return jsonify({"success": True})


# ── Tags ──────────────────────────────────────────────────────────────────────

@app.route("/tags")
def get_tags():
    with get_db() as conn:
        return jsonify([dict(r) for r in conn.execute("SELECT * FROM tags ORDER BY name").fetchall()])


@app.route("/tags", methods=["POST"])
def create_tag():
    data  = request.get_json() or {}
    name  = data.get("name", "").strip().lower()
    color = data.get("color", "#3b82f6")
    if not name:
        return jsonify({"error": "Name required"}), 400
    with get_db() as conn:
        try:
            cur = conn.execute("INSERT INTO tags (name,color) VALUES (?,?)", (name, color))
            conn.commit()
            return jsonify(dict(conn.execute("SELECT * FROM tags WHERE id=?", (cur.lastrowid,)).fetchone())), 201
        except Exception:
            return jsonify(dict(conn.execute("SELECT * FROM tags WHERE name=?", (name,)).fetchone()))


@app.route("/tags/<int:tag_id>", methods=["DELETE"])
def delete_tag(tag_id):
    with get_db() as conn:
        conn.execute("DELETE FROM tags WHERE id=?", (tag_id,))
        conn.commit()
        return jsonify({"success": True})


@app.route("/tasks/<int:tid>/tags/<int:tag_id>", methods=["POST"])
def attach_tag(tid, tag_id):
    with get_db() as conn:
        try:
            conn.execute("INSERT INTO task_tags (task_id,tag_id) VALUES (?,?)", (tid, tag_id))
            conn.commit()
        except Exception:
            pass
        return jsonify({"success": True})


@app.route("/tasks/<int:tid>/tags/<int:tag_id>", methods=["DELETE"])
def detach_tag(tid, tag_id):
    with get_db() as conn:
        conn.execute("DELETE FROM task_tags WHERE task_id=? AND tag_id=?", (tid, tag_id))
        conn.commit()
        return jsonify({"success": True})


# ── Task dependencies ────────────────────────────────────────────────────────

@app.route("/tasks/<int:tid>/dependencies", methods=["POST"])
def add_dependency(tid):
    data   = request.get_json() or {}
    dep_id = data.get("depends_on_id")
    if not dep_id or int(dep_id) == tid:
        return jsonify({"error": "Invalid dependency"}), 400
    dep_id = int(dep_id)
    with get_db() as conn:
        if not conn.execute("SELECT id FROM tasks WHERE id=?", (tid,)).fetchone():
            return jsonify({"error": "Task not found"}), 404
        if not conn.execute("SELECT id FROM tasks WHERE id=?", (dep_id,)).fetchone():
            return jsonify({"error": "Dependency task not found"}), 404
        if _has_dependency_path(conn, dep_id, tid):
            return jsonify({"error": "Would create a circular dependency"}), 400
        try:
            conn.execute(
                "INSERT INTO task_dependencies (task_id,depends_on_id,created_at) VALUES (?,?,?)",
                (tid, dep_id, _now()),
            )
            dep_task = conn.execute("SELECT title FROM tasks WHERE id=?", (dep_id,)).fetchone()
            _log(conn, tid, "dependency_added", f"now depends on \"{dep_task['title']}\"")
            conn.commit()
        except sqlite3.IntegrityError:
            pass
        return jsonify({"success": True}), 201


@app.route("/tasks/<int:tid>/dependencies/<int:dep_id>", methods=["DELETE"])
def remove_dependency(tid, dep_id):
    with get_db() as conn:
        conn.execute(
            "DELETE FROM task_dependencies WHERE task_id=? AND depends_on_id=?", (tid, dep_id)
        )
        _log(conn, tid, "dependency_removed", f"no longer depends on #{dep_id}")
        conn.commit()
        return jsonify({"success": True})


# ── Notes (Books › Chapters › Pages) ────────────────────────────────────────────

def _enrich_notebooks(conn, book_rows):
    result = []
    for b in book_rows:
        bd = dict(b)
        chapters = []
        for c in conn.execute(
            "SELECT * FROM note_chapters WHERE book_id=? ORDER BY position, id", (bd["id"],)
        ).fetchall():
            cd = dict(c)
            cd["pages"] = [dict(p) for p in conn.execute(
                "SELECT * FROM note_pages WHERE chapter_id=? ORDER BY position, id", (cd["id"],)
            ).fetchall()]
            chapters.append(cd)
        bd["chapters"] = chapters
        result.append(bd)
    return result


@app.route("/notes/books")
def get_notebooks():
    with get_db() as conn:
        rows = conn.execute("SELECT * FROM note_books ORDER BY position, id").fetchall()
        return jsonify(_enrich_notebooks(conn, rows))


@app.route("/notes/books", methods=["POST"])
def create_notebook():
    data  = request.get_json() or {}
    title = data.get("title", "").strip()
    if not title:
        return jsonify({"error": "Title required"}), 400
    now = _now()
    with get_db() as conn:
        pos = conn.execute("SELECT COALESCE(MAX(position),0)+1 FROM note_books").fetchone()[0]
        cur = conn.execute(
            "INSERT INTO note_books (title,color,position,created_at,updated_at) VALUES (?,?,?,?,?)",
            (title, data.get("color", "default") or "default", pos, now, now),
        )
        conn.commit()
        row = conn.execute("SELECT * FROM note_books WHERE id=?", (cur.lastrowid,)).fetchone()
        bd = dict(row); bd["chapters"] = []
        return jsonify(bd), 201


@app.route("/notes/books/<int:bid>", methods=["PUT"])
def update_notebook(bid):
    data = request.get_json() or {}
    with get_db() as conn:
        b = conn.execute("SELECT * FROM note_books WHERE id=?", (bid,)).fetchone()
        if not b:
            return jsonify({"error": "Not found"}), 404
        bd  = dict(b)
        now = _now()
        conn.execute(
            "UPDATE note_books SET title=?,color=?,updated_at=? WHERE id=?",
            (data.get("title", bd["title"]).strip() or bd["title"],
             data.get("color", bd["color"]) or "default", now, bid),
        )
        conn.commit()
        row = conn.execute("SELECT * FROM note_books WHERE id=?", (bid,)).fetchone()
        return jsonify(dict(row))


@app.route("/notes/books/<int:bid>", methods=["DELETE"])
def delete_notebook(bid):
    with get_db() as conn:
        if not conn.execute("SELECT id FROM note_books WHERE id=?", (bid,)).fetchone():
            return jsonify({"error": "Not found"}), 404
        conn.execute("DELETE FROM note_books WHERE id=?", (bid,))
        conn.commit()
        return jsonify({"success": True})


@app.route("/notes/books/<int:bid>/chapters", methods=["POST"])
def create_chapter(bid):
    data  = request.get_json() or {}
    title = data.get("title", "").strip()
    if not title:
        return jsonify({"error": "Title required"}), 400
    now = _now()
    with get_db() as conn:
        if not conn.execute("SELECT id FROM note_books WHERE id=?", (bid,)).fetchone():
            return jsonify({"error": "Notebook not found"}), 404
        pos = conn.execute(
            "SELECT COALESCE(MAX(position),0)+1 FROM note_chapters WHERE book_id=?", (bid,)
        ).fetchone()[0]
        cur = conn.execute(
            "INSERT INTO note_chapters (book_id,title,position,created_at,updated_at) VALUES (?,?,?,?,?)",
            (bid, title, pos, now, now),
        )
        conn.execute("UPDATE note_books SET updated_at=? WHERE id=?", (now, bid))
        conn.commit()
        row = conn.execute("SELECT * FROM note_chapters WHERE id=?", (cur.lastrowid,)).fetchone()
        cd = dict(row); cd["pages"] = []
        return jsonify(cd), 201


@app.route("/notes/chapters/<int:cid>", methods=["PUT"])
def update_chapter(cid):
    data = request.get_json() or {}
    with get_db() as conn:
        c = conn.execute("SELECT * FROM note_chapters WHERE id=?", (cid,)).fetchone()
        if not c:
            return jsonify({"error": "Not found"}), 404
        cd  = dict(c)
        now = _now()
        conn.execute(
            "UPDATE note_chapters SET title=?,updated_at=? WHERE id=?",
            (data.get("title", cd["title"]).strip() or cd["title"], now, cid),
        )
        conn.execute("UPDATE note_books SET updated_at=? WHERE id=?", (now, cd["book_id"]))
        conn.commit()
        row = conn.execute("SELECT * FROM note_chapters WHERE id=?", (cid,)).fetchone()
        return jsonify(dict(row))


@app.route("/notes/chapters/<int:cid>", methods=["DELETE"])
def delete_chapter(cid):
    with get_db() as conn:
        c = conn.execute("SELECT * FROM note_chapters WHERE id=?", (cid,)).fetchone()
        if not c:
            return jsonify({"error": "Not found"}), 404
        conn.execute("DELETE FROM note_chapters WHERE id=?", (cid,))
        conn.execute("UPDATE note_books SET updated_at=? WHERE id=?", (_now(), c["book_id"]))
        conn.commit()
        return jsonify({"success": True})


@app.route("/notes/chapters/<int:cid>/pages", methods=["POST"])
def create_page(cid):
    data  = request.get_json() or {}
    title = data.get("title", "").strip()
    if not title:
        return jsonify({"error": "Title required"}), 400
    now = _now()
    with get_db() as conn:
        chapter = conn.execute("SELECT * FROM note_chapters WHERE id=?", (cid,)).fetchone()
        if not chapter:
            return jsonify({"error": "Chapter not found"}), 404
        pos = conn.execute(
            "SELECT COALESCE(MIN(position),0)-1 FROM note_pages WHERE chapter_id=?", (cid,)
        ).fetchone()[0]
        cur = conn.execute(
            "INSERT INTO note_pages (chapter_id,title,body,color,pinned,position,created_at,updated_at)"
            " VALUES (?,?,?,?,?,?,?,?)",
            (cid, title, data.get("body", "") or "", data.get("color", "default") or "default",
             1 if data.get("pinned") else 0, pos, now, now),
        )
        conn.execute("UPDATE note_chapters SET updated_at=? WHERE id=?", (now, cid))
        conn.execute("UPDATE note_books SET updated_at=? WHERE id=?", (now, chapter["book_id"]))
        conn.commit()
        row = conn.execute("SELECT * FROM note_pages WHERE id=?", (cur.lastrowid,)).fetchone()
        return jsonify(dict(row)), 201


@app.route("/notes/pages/<int:pid>", methods=["PUT"])
def update_page(pid):
    data = request.get_json() or {}
    with get_db() as conn:
        p = conn.execute("SELECT * FROM note_pages WHERE id=?", (pid,)).fetchone()
        if not p:
            return jsonify({"error": "Not found"}), 404
        pd  = dict(p)
        now = _now()
        conn.execute(
            "UPDATE note_pages SET title=?,body=?,color=?,pinned=?,updated_at=? WHERE id=?",
            (data.get("title", pd["title"]), data.get("body", pd["body"]),
             data.get("color", pd["color"]) or "default",
             1 if data.get("pinned", pd["pinned"]) else 0, now, pid),
        )
        chapter = conn.execute("SELECT * FROM note_chapters WHERE id=?", (pd["chapter_id"],)).fetchone()
        conn.execute("UPDATE note_chapters SET updated_at=? WHERE id=?", (now, pd["chapter_id"]))
        if chapter:
            conn.execute("UPDATE note_books SET updated_at=? WHERE id=?", (now, chapter["book_id"]))
        conn.commit()
        row = conn.execute("SELECT * FROM note_pages WHERE id=?", (pid,)).fetchone()
        return jsonify(dict(row))


@app.route("/notes/pages/<int:pid>", methods=["DELETE"])
def delete_page(pid):
    with get_db() as conn:
        if not conn.execute("SELECT id FROM note_pages WHERE id=?", (pid,)).fetchone():
            return jsonify({"error": "Not found"}), 404
        conn.execute("DELETE FROM note_pages WHERE id=?", (pid,))
        conn.commit()
        return jsonify({"success": True})


# ── Activity log ──────────────────────────────────────────────────────────────

@app.route("/tasks/<int:tid>/activity")
def get_activity(tid):
    with get_db() as conn:
        rows = conn.execute(
            "SELECT * FROM activity_log WHERE task_id=? ORDER BY created_at DESC LIMIT 50", (tid,)
        ).fetchall()
        return jsonify([dict(r) for r in rows])


# ── Time tracking ─────────────────────────────────────────────────────────────

@app.route("/tasks/<int:tid>/time/start", methods=["POST"])
def start_timer(tid):
    now = _now()
    with get_db() as conn:
        active = conn.execute(
            "SELECT * FROM time_entries WHERE task_id=? AND ended_at IS NULL", (tid,)
        ).fetchone()
        if active:
            dur = int((datetime.fromisoformat(now) - datetime.fromisoformat(active["started_at"])).total_seconds())
            conn.execute("UPDATE time_entries SET ended_at=?,duration_seconds=? WHERE id=?",
                         (now, dur, active["id"]))
        cur = conn.execute("INSERT INTO time_entries (task_id,started_at) VALUES (?,?)", (tid, now))
        _log(conn, tid, "timer_started")
        conn.commit()
        return jsonify(dict(conn.execute("SELECT * FROM time_entries WHERE id=?", (cur.lastrowid,)).fetchone())), 201


@app.route("/tasks/<int:tid>/time/stop", methods=["POST"])
def stop_timer(tid):
    now = _now()
    with get_db() as conn:
        active = conn.execute(
            "SELECT * FROM time_entries WHERE task_id=? AND ended_at IS NULL", (tid,)
        ).fetchone()
        if not active:
            return jsonify({"error": "No active timer"}), 400
        dur = int((datetime.fromisoformat(now) - datetime.fromisoformat(active["started_at"])).total_seconds())
        conn.execute("UPDATE time_entries SET ended_at=?,duration_seconds=? WHERE id=?",
                     (now, dur, active["id"]))
        _log(conn, tid, "timer_stopped", f"{dur}s")
        conn.commit()
        return jsonify({"duration_seconds": dur})


@app.route("/tasks/<int:tid>/time")
def get_time(tid):
    with get_db() as conn:
        rows  = conn.execute(
            "SELECT * FROM time_entries WHERE task_id=? ORDER BY started_at DESC", (tid,)
        ).fetchall()
        total = sum(r["duration_seconds"] or 0 for r in rows)
        return jsonify({"entries": [dict(r) for r in rows], "total_seconds": total})


# ── Reminders ─────────────────────────────────────────────────────────────────

@app.route("/reminders")
def get_all_reminders():
    with get_db() as conn:
        return jsonify([dict(r) for r in
                        conn.execute("SELECT * FROM reminders WHERE is_active=1").fetchall()])


@app.route("/reminders/due")
def get_due_reminders():
    now = _now()
    with get_db() as conn:
        task_rows = conn.execute(
            "SELECT r.*, t.title AS task_title, 'task' AS rem_type FROM reminders r "
            "JOIN tasks t ON r.task_id=t.id "
            "WHERE r.is_active=1 AND r.remind_at<=? "
            "AND (r.snoozed_until IS NULL OR r.snoozed_until<=?) "
            "AND t.status NOT IN ('completed','cancelled')",
            (now, now),
        ).fetchall()
        standalone_rows = conn.execute(
            "SELECT *, 'standalone' AS rem_type FROM standalone_reminders "
            "WHERE is_active=1 AND remind_at<=? "
            "AND (snoozed_until IS NULL OR snoozed_until<=?)",
            (now, now),
        ).fetchall()
        return jsonify([dict(r) for r in task_rows] + [dict(r) for r in standalone_rows])


@app.route("/tasks/<int:tid>/reminder")
def get_task_reminder(tid):
    with get_db() as conn:
        row = conn.execute(
            "SELECT * FROM reminders WHERE task_id=? AND is_active=1 ORDER BY id DESC LIMIT 1", (tid,)
        ).fetchone()
        return jsonify(dict(row) if row else {})


@app.route("/tasks/<int:tid>/reminder", methods=["POST"])
def set_reminder(tid):
    data      = request.get_json() or {}
    remind_at = data.get("remind_at", "").strip()
    if not remind_at:
        return jsonify({"error": "remind_at required"}), 400
    sound = 1 if data.get("sound", True) else 0
    with get_db() as conn:
        conn.execute("UPDATE reminders SET is_active=0 WHERE task_id=?", (tid,))
        cur = conn.execute(
            "INSERT INTO reminders (task_id,remind_at,recurrence,sound,is_active) VALUES (?,?,?,?,1)",
            (tid, remind_at, data.get("recurrence","none"), sound),
        )
        conn.commit()
        return jsonify(dict(conn.execute("SELECT * FROM reminders WHERE id=?", (cur.lastrowid,)).fetchone())), 201


@app.route("/tasks/<int:tid>/reminder", methods=["DELETE"])
def remove_reminder(tid):
    with get_db() as conn:
        conn.execute("UPDATE reminders SET is_active=0 WHERE task_id=?", (tid,))
        conn.commit()
        return jsonify({"success": True})


@app.route("/reminders/<int:rid>/acknowledge", methods=["POST"])
def acknowledge_reminder(rid):
    now = _now()
    data = request.get_json() or {}
    snooze_minutes = data.get("snooze_minutes")
    with get_db() as conn:
        rem = conn.execute("SELECT * FROM reminders WHERE id=?", (rid,)).fetchone()
        if not rem:
            return jsonify({"error": "Not found"}), 404
        if snooze_minutes:
            snoozed_until = (datetime.now() + timedelta(minutes=int(snooze_minutes))).isoformat(timespec="seconds")
            conn.execute("UPDATE reminders SET snoozed_until=?,last_triggered=? WHERE id=?",
                         (snoozed_until, now, rid))
        elif rem["recurrence"] == "none":
            conn.execute("UPDATE reminders SET is_active=0,last_triggered=?,snoozed_until=NULL WHERE id=?",
                         (now, rid))
        else:
            nxt = _next_occurrence(rem["remind_at"], rem["recurrence"])
            conn.execute("UPDATE reminders SET remind_at=?,last_triggered=?,snoozed_until=NULL WHERE id=?",
                         (nxt, now, rid))
        conn.commit()
        return jsonify({"success": True})


# ── Standalone Reminders ──────────────────────────────────────────────────────

@app.route("/standalone-reminders")
def get_standalone_reminders():
    with get_db() as conn:
        rows = conn.execute(
            "SELECT * FROM standalone_reminders WHERE is_active=1 ORDER BY remind_at"
        ).fetchall()
        return jsonify([dict(r) for r in rows])


@app.route("/standalone-reminders", methods=["POST"])
def create_standalone_reminder():
    data = request.get_json() or {}
    title = data.get("title", "").strip()
    if not title:
        return jsonify({"error": "Title required"}), 400
    remind_at = data.get("remind_at", "").strip()
    if not remind_at:
        return jsonify({"error": "remind_at required"}), 400
    now = _now()
    with get_db() as conn:
        cur = conn.execute(
            "INSERT INTO standalone_reminders "
            "(title,note,reminder_type,remind_at,recurrence,sound,is_active,created_at) "
            "VALUES (?,?,?,?,?,?,1,?)",
            (title, data.get("note", ""), data.get("reminder_type", "general"),
             remind_at, data.get("recurrence", "none"),
             1 if data.get("sound", True) else 0, now),
        )
        conn.commit()
        row = conn.execute("SELECT * FROM standalone_reminders WHERE id=?", (cur.lastrowid,)).fetchone()
        return jsonify(dict(row)), 201


@app.route("/standalone-reminders/<int:rid>", methods=["PUT"])
def update_standalone_reminder(rid):
    data = request.get_json() or {}
    with get_db() as conn:
        rem = conn.execute("SELECT * FROM standalone_reminders WHERE id=?", (rid,)).fetchone()
        if not rem:
            return jsonify({"error": "Not found"}), 404
        conn.execute(
            "UPDATE standalone_reminders SET title=?,note=?,reminder_type=?,remind_at=?,recurrence=?,sound=? WHERE id=?",
            (data.get("title", rem["title"]).strip() or rem["title"],
             data.get("note", rem["note"]),
             data.get("reminder_type", rem["reminder_type"]),
             data.get("remind_at", rem["remind_at"]),
             data.get("recurrence", rem["recurrence"]),
             1 if data.get("sound", rem["sound"]) else 0, rid),
        )
        conn.commit()
        return jsonify(dict(conn.execute("SELECT * FROM standalone_reminders WHERE id=?", (rid,)).fetchone()))


@app.route("/standalone-reminders/<int:rid>", methods=["DELETE"])
def delete_standalone_reminder(rid):
    with get_db() as conn:
        conn.execute("DELETE FROM standalone_reminders WHERE id=?", (rid,))
        conn.commit()
        return jsonify({"success": True})


@app.route("/standalone-reminders/<int:rid>/acknowledge", methods=["POST"])
def acknowledge_standalone_reminder(rid):
    now = _now()
    data = request.get_json() or {}
    snooze_minutes = data.get("snooze_minutes")
    with get_db() as conn:
        rem = conn.execute("SELECT * FROM standalone_reminders WHERE id=?", (rid,)).fetchone()
        if not rem:
            return jsonify({"error": "Not found"}), 404
        if snooze_minutes:
            snoozed_until = (datetime.now() + timedelta(minutes=int(snooze_minutes))).isoformat(timespec="seconds")
            conn.execute("UPDATE standalone_reminders SET snoozed_until=?,last_triggered=? WHERE id=?",
                         (snoozed_until, now, rid))
        elif rem["recurrence"] == "none":
            conn.execute("UPDATE standalone_reminders SET is_active=0,last_triggered=?,snoozed_until=NULL WHERE id=?",
                         (now, rid))
        else:
            nxt = _next_occurrence(rem["remind_at"], rem["recurrence"])
            conn.execute("UPDATE standalone_reminders SET remind_at=?,last_triggered=?,snoozed_until=NULL WHERE id=?",
                         (nxt, now, rid))
        conn.commit()
        return jsonify({"success": True})


# ── Meetings ──────────────────────────────────────────────────────────────────

@app.route("/meetings")
def get_meetings():
    with get_db() as conn:
        rows = conn.execute(
            "SELECT * FROM meetings ORDER BY date, time"
        ).fetchall()
        return jsonify(_enrich_meetings(conn, rows))


@app.route("/meetings", methods=["POST"])
def create_meeting():
    data  = request.get_json() or {}
    title = data.get("title", "").strip()
    if not title:
        return jsonify({"error": "Title required"}), 400
    now = _now()
    with get_db() as conn:
        cur = conn.execute(
            "INSERT INTO meetings (title,description,date,time,duration_minutes,location,recurrence,status,created_at,updated_at)"
            " VALUES (?,?,?,?,?,?,?,?,?,?)",
            (title, data.get("description", "").strip(),
             data.get("date", ""), data.get("time") or None,
             int(data.get("duration_minutes", 60) or 60),
             data.get("location", "").strip(),
             data.get("recurrence", "none") or "none",
             data.get("status", "upcoming") or "upcoming", now, now),
        )
        mid = cur.lastrowid
        conn.commit()
        row = conn.execute("SELECT * FROM meetings WHERE id=?", (mid,)).fetchone()
        return jsonify(_enrich_meetings(conn, [row])[0]), 201


@app.route("/meetings/<int:mid>", methods=["PUT"])
def update_meeting(mid):
    data = request.get_json() or {}
    with get_db() as conn:
        m = conn.execute("SELECT * FROM meetings WHERE id=?", (mid,)).fetchone()
        if not m:
            return jsonify({"error": "Not found"}), 404
        md  = dict(m)
        now = _now()
        new_status     = data.get("status",     md["status"])
        new_recurrence = data.get("recurrence", md["recurrence"]) or "none"
        conn.execute(
            "UPDATE meetings SET title=?,description=?,date=?,time=?,duration_minutes=?,location=?,recurrence=?,status=?,updated_at=? WHERE id=?",
            (data.get("title", md["title"]).strip() or md["title"],
             data.get("description", md["description"]),
             data.get("date",     md["date"]),
             data.get("time",     md["time"]) or None,
             int(data.get("duration_minutes", md["duration_minutes"]) or 60),
             data.get("location", md["location"]),
             new_recurrence, new_status, now, mid),
        )
        if new_status == "completed" and md["status"] != "completed" and new_recurrence != "none":
            updated_m = conn.execute("SELECT * FROM meetings WHERE id=?", (mid,)).fetchone()
            _spawn_next_meeting(conn, dict(updated_m), now)
        conn.commit()
        updated = conn.execute("SELECT * FROM meetings WHERE id=?", (mid,)).fetchone()
        return jsonify(_enrich_meetings(conn, [updated])[0])


@app.route("/meetings/<int:mid>", methods=["DELETE"])
def delete_meeting(mid):
    with get_db() as conn:
        if not conn.execute("SELECT id FROM meetings WHERE id=?", (mid,)).fetchone():
            return jsonify({"error": "Not found"}), 404
        conn.execute("DELETE FROM meetings WHERE id=?", (mid,))
        conn.commit()
        return jsonify({"success": True})


@app.route("/meetings/<int:mid>/points")
def get_meeting_points(mid):
    with get_db() as conn:
        rows = conn.execute(
            "SELECT * FROM meeting_points WHERE meeting_id=? ORDER BY position, id", (mid,)
        ).fetchall()
        return jsonify([dict(r) for r in rows])


@app.route("/meetings/<int:mid>/points", methods=["POST"])
def add_meeting_point(mid):
    data    = request.get_json() or {}
    content = data.get("content", "").strip()
    if not content:
        return jsonify({"error": "Content required"}), 400
    now = _now()
    with get_db() as conn:
        pos = conn.execute(
            "SELECT COALESCE(MAX(position),0)+1 FROM meeting_points WHERE meeting_id=?", (mid,)
        ).fetchone()[0]
        cur = conn.execute(
            "INSERT INTO meeting_points (meeting_id,content,is_done,position,created_at) VALUES (?,?,0,?,?)",
            (mid, content, pos, now),
        )
        conn.commit()
        row = conn.execute("SELECT * FROM meeting_points WHERE id=?", (cur.lastrowid,)).fetchone()
        return jsonify(dict(row)), 201


@app.route("/meeting-points/<int:pid>", methods=["PUT"])
def update_meeting_point(pid):
    data = request.get_json() or {}
    with get_db() as conn:
        p = conn.execute("SELECT * FROM meeting_points WHERE id=?", (pid,)).fetchone()
        if not p:
            return jsonify({"error": "Not found"}), 404
        conn.execute(
            "UPDATE meeting_points SET content=?,is_done=? WHERE id=?",
            (data.get("content", p["content"]), data.get("is_done", p["is_done"]), pid),
        )
        conn.commit()
        return jsonify(dict(conn.execute("SELECT * FROM meeting_points WHERE id=?", (pid,)).fetchone()))


@app.route("/meeting-points/<int:pid>", methods=["DELETE"])
def delete_meeting_point(pid):
    with get_db() as conn:
        conn.execute("DELETE FROM meeting_points WHERE id=?", (pid,))
        conn.commit()
        return jsonify({"success": True})


# ── Task Steps (workflow phases, each holding real tasks) ─────────────────────

def _enrich_steps(conn, tid):
    steps = conn.execute(
        "SELECT * FROM task_steps WHERE task_id=? ORDER BY position, id", (tid,)
    ).fetchall()
    result = []
    for s in steps:
        sd = dict(s)
        step_tasks = conn.execute(
            "SELECT * FROM tasks WHERE step_id=? ORDER BY created_at", (s["id"],)
        ).fetchall()
        enriched = _enrich(conn, step_tasks)
        sd["tasks"]      = enriched
        sd["task_total"] = len(enriched)
        sd["task_done"]  = sum(1 for t in enriched if t["status"] == "completed")
        result.append(sd)
    return result


@app.route("/tasks/<int:tid>/steps")
def get_steps(tid):
    with get_db() as conn:
        return jsonify(_enrich_steps(conn, tid))


@app.route("/tasks/<int:tid>/steps", methods=["POST"])
def add_step(tid):
    data  = request.get_json() or {}
    title = data.get("title", "").strip()
    if not title:
        return jsonify({"error": "Title required"}), 400
    now = _now()
    with get_db() as conn:
        default_pos = conn.execute(
            "SELECT COALESCE(MAX(position),0)+1 FROM task_steps WHERE task_id=?", (tid,)
        ).fetchone()[0]
        pos = int(data.get("position") or default_pos)
        cur = conn.execute(
            "INSERT INTO task_steps (task_id,title,status,position,parallel_group,created_at,updated_at) VALUES (?,?,?,?,?,?,?)",
            (tid, title, "not_started", pos, data.get("parallel_group") or None, now, now),
        )
        conn.commit()
        sd = dict(conn.execute("SELECT * FROM task_steps WHERE id=?", (cur.lastrowid,)).fetchone())
        sd["tasks"] = []
        sd["task_total"] = 0
        sd["task_done"]  = 0
        return jsonify(sd), 201


@app.route("/steps/<int:sid>")
def get_step(sid):
    with get_db() as conn:
        s = conn.execute("SELECT * FROM task_steps WHERE id=?", (sid,)).fetchone()
        if not s:
            return jsonify({"error": "Not found"}), 404
        sd = dict(s)
        step_tasks = conn.execute("SELECT * FROM tasks WHERE step_id=? ORDER BY created_at", (sid,)).fetchall()
        enriched = _enrich(conn, step_tasks)
        sd["tasks"]      = enriched
        sd["task_total"] = len(enriched)
        sd["task_done"]  = sum(1 for t in enriched if t["status"] == "completed")
        return jsonify(sd)


@app.route("/steps/<int:sid>", methods=["PUT"])
def update_step(sid):
    data = request.get_json() or {}
    with get_db() as conn:
        s = conn.execute("SELECT * FROM task_steps WHERE id=?", (sid,)).fetchone()
        if not s:
            return jsonify({"error": "Not found"}), 404
        now = _now()
        pg = data["parallel_group"] if "parallel_group" in data else dict(s).get("parallel_group")
        conn.execute(
            "UPDATE task_steps SET title=?,status=?,parallel_group=?,updated_at=? WHERE id=?",
            (data.get("title", s["title"]).strip() or s["title"],
             data.get("status", s["status"]), pg, now, sid),
        )
        conn.commit()
        sd = dict(conn.execute("SELECT * FROM task_steps WHERE id=?", (sid,)).fetchone())
        step_tasks = conn.execute("SELECT * FROM tasks WHERE step_id=? ORDER BY created_at", (sid,)).fetchall()
        enriched = _enrich(conn, step_tasks)
        sd["tasks"]      = enriched
        sd["task_total"] = len(enriched)
        sd["task_done"]  = sum(1 for t in enriched if t["status"] == "completed")
        return jsonify(sd)


@app.route("/steps/<int:sid>", methods=["DELETE"])
def delete_step(sid):
    with get_db() as conn:
        if not conn.execute("SELECT id FROM task_steps WHERE id=?", (sid,)).fetchone():
            return jsonify({"error": "Not found"}), 404
        conn.execute("UPDATE tasks SET step_id=NULL WHERE step_id=?", (sid,))
        conn.execute("DELETE FROM task_steps WHERE id=?", (sid,))
        conn.commit()
        return jsonify({"success": True})


# ── Excel export (multi-sheet) ────────────────────────────────────────────────

def _xlsx_header(ws, cols, color="1e3a5f"):
    fill = PatternFill("solid", fgColor=color)
    font = Font(bold=True, color="FFFFFF", size=10)
    border = Border(
        bottom=Side(style="thin", color="334155"),
        right=Side(style="thin",  color="334155"),
    )
    for i, col in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i, value=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="center", wrap_text=False)
        cell.border = border
    ws.row_dimensions[1].height = 18

def _xlsx_row(ws, row_num, values, alt=False):
    fill = PatternFill("solid", fgColor="111827" if alt else "0a0f1e")
    for i, val in enumerate(values, 1):
        cell = ws.cell(row=row_num, column=i, value=val)
        cell.font = Font(size=9, color="e2e8f0")
        cell.alignment = Alignment(vertical="center")
        cell.fill = fill

def _auto_width(ws, min_w=8, max_w=50):
    for col in ws.columns:
        length = max(
            min_w,
            min(max_w, max((len(str(c.value or "")) for c in col), default=min_w))
        )
        ws.column_dimensions[get_column_letter(col[0].column)].width = length + 2


@app.route("/export.xlsx")
def export_xlsx():
    type_filter = request.args.get("type", "all")
    with get_db() as conn:
        # Tasks
        if type_filter and type_filter != "all":
            task_rows = conn.execute(
                "SELECT * FROM tasks WHERE task_type=? ORDER BY created_at", (type_filter,)
            ).fetchall()
        else:
            task_rows = conn.execute("SELECT * FROM tasks ORDER BY created_at").fetchall()
        tasks = _enrich(conn, task_rows)

        # All subtasks
        all_subs = conn.execute(
            "SELECT s.*, t.title AS task_title FROM subtasks s JOIN tasks t ON s.task_id=t.id ORDER BY s.task_id, s.position"
        ).fetchall()

        # Meetings
        meeting_rows = conn.execute("SELECT * FROM meetings ORDER BY date, time").fetchall()
        meetings = _enrich_meetings(conn, meeting_rows)

        # All discussion points
        all_points = conn.execute(
            "SELECT mp.*, m.title AS meeting_title, m.date AS meeting_date "
            "FROM meeting_points mp JOIN meetings m ON mp.meeting_id=m.id ORDER BY mp.meeting_id, mp.position"
        ).fetchall()

        # Workflow steps, with task progress rolled up
        all_steps = conn.execute(
            "SELECT ts.*, t.title AS task_title,"
            " (SELECT COUNT(*) FROM tasks tk WHERE tk.step_id=ts.id) AS task_total,"
            " (SELECT COUNT(*) FROM tasks tk WHERE tk.step_id=ts.id AND tk.status='completed') AS task_done"
            " FROM task_steps ts JOIN tasks t ON ts.task_id=t.id ORDER BY ts.task_id, ts.position"
        ).fetchall()

    wb = Workbook()

    # ── Sheet 1: Tasks ────────────────────────────────────────────────────────
    ws_tasks = wb.active
    ws_tasks.title = "Tasks"
    task_cols = ["ID","Title","Type","Status","Priority","Due Date","Due Time",
                 "Est (min)","Project","Workflow","Step","Blocked On","Recurrence","Tags","Subtasks","Time (hrs)","Created","Updated","Description"]
    _xlsx_header(ws_tasks, task_cols)
    for i, t in enumerate(tasks, 2):
        sub_titles = "; ".join(
            s["title"] for s in all_subs
            if s["task_id"] == t["id"]
        ) or ""
        blocked_on = "; ".join(
            d["title"] for d in t["depends_on"] if d["status"] not in ("completed", "cancelled")
        )
        step = t.get("step")
        _xlsx_row(ws_tasks, i, [
            t["id"], t["title"], dict(t).get("task_type","general"),
            t["status"], t["priority"],
            t["due_date"] or "", t["due_time"] or "",
            t["estimated_minutes"] or 0, t["project"] or "",
            step["workflow_title"] if step else "", step["title"] if step else "",
            blocked_on,
            t["recurrence"] or "none",
            ", ".join(tg["name"] for tg in t["tags"]),
            sub_titles,
            round(t["time_total"] / 3600, 2),
            t["created_at"], t["updated_at"], t["description"],
        ], alt=i % 2 == 0)
    _auto_width(ws_tasks)
    ws_tasks.freeze_panes = "A2"

    # ── Sheet 2: Subtasks ─────────────────────────────────────────────────────
    ws_subs = wb.create_sheet("Subtasks")
    _xlsx_header(ws_subs, ["Task ID","Task Title","Subtask","Completed"], color="1a3550")
    for i, s in enumerate(all_subs, 2):
        _xlsx_row(ws_subs, i, [
            s["task_id"], s["task_title"], s["title"],
            "Yes" if s["is_done"] else "No",
        ], alt=i % 2 == 0)
    _auto_width(ws_subs)
    ws_subs.freeze_panes = "A2"

    # ── Sheet 3: Meetings ─────────────────────────────────────────────────────
    ws_mtg = wb.create_sheet("Meetings")
    mtg_cols = ["ID","Title","Date","Time","Duration (min)","Location",
                "Recurrence","Status","Points Done","Points Total","Points List","Notes","Created"]
    _xlsx_header(ws_mtg, mtg_cols, color="1e3a2a")
    for i, m in enumerate(meetings, 2):
        pts_all = [p for p in all_points if p["meeting_id"] == m["id"]]
        pts_str = "; ".join(p["content"] for p in pts_all)
        _xlsx_row(ws_mtg, i, [
            m["id"], m["title"], m["date"], m["time"] or "",
            m["duration_minutes"], m["location"] or "",
            m["recurrence"] or "none", m["status"],
            m["point_done"], m["point_total"], pts_str,
            m["description"] or "", m["created_at"],
        ], alt=i % 2 == 0)
    _auto_width(ws_mtg)
    ws_mtg.freeze_panes = "A2"

    # ── Sheet 4: Discussion Points ────────────────────────────────────────────
    ws_pts = wb.create_sheet("Discussion Points")
    _xlsx_header(ws_pts, ["Meeting ID","Meeting Title","Meeting Date","Point","Done"], color="2a1e3a")
    for i, p in enumerate(all_points, 2):
        _xlsx_row(ws_pts, i, [
            p["meeting_id"], p["meeting_title"], p["meeting_date"],
            p["content"], "Yes" if p["is_done"] else "No",
        ], alt=i % 2 == 0)
    _auto_width(ws_pts)
    ws_pts.freeze_panes = "A2"

    # ── Sheet 5: Workflow Steps ────────────────────────────────────────────────
    ws_steps = wb.create_sheet("Workflow Steps")
    _xlsx_header(ws_steps, ["Workflow ID","Workflow Title","Step #","Step Title","Status","Tasks Done","Tasks Total","Created"], color="1a2a3a")
    for i, s in enumerate(all_steps, 2):
        _xlsx_row(ws_steps, i, [
            s["task_id"], s["task_title"], s["position"],
            s["title"], s["status"], s["task_done"], s["task_total"], s["created_at"],
        ], alt=i % 2 == 0)
    _auto_width(ws_steps)
    ws_steps.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"taskflow_{type_filter}.xlsx" if type_filter != "all" else "taskflow_export.xlsx"
    return Response(
        buf.read(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={fname}"},
    )


init_db()

if __name__ == "__main__":
    app.run(debug=True, port=5001)
